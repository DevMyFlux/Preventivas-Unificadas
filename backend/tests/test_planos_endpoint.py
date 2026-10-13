"""
Regressão: a rota /planos é cadastral (seção 7 do briefing original) e precisa mostrar planos
ativos E inativos — cada item já vem com o campo "ativo" para a tela exibir o status. Bug
antigo: brasilandia/routes.py chamava filtros_planos() sem passar ativo=None, caindo no default
(ativo=True) e escondendo os inativos. Agora as duas unidades usam a MESMA implementação
(core/unit_blueprint.py); o teste garante que continuam assim.
"""
import pytest

from tests.conftest import fazer_login


@pytest.fixture()
def capturar_filtros(monkeypatch):
    from core import unit_blueprint as ub
    capturado = {}

    def fake_paginar(headers, payload, endpoint):
        capturado[endpoint] = payload["filterGroups"][0]["filters"]
        return []

    monkeypatch.setattr(ub, "paginar", fake_paginar)
    monkeypatch.setattr(ub, "get_headers", lambda: {})
    return capturado


@pytest.mark.parametrize("login,prefixo", [("ana.hmb", "/api/brasilandia"), ("beto.hetrin", "/api/grandmassif")])
def test_planos_nao_filtra_por_ativo(client, criar_usuarios, capturar_filtros, login, prefixo):
    fazer_login(client, login, criar_usuarios)
    assert client.get(f"{prefixo}/planos").status_code == 200
    propriedades = [f["property"] for f in capturar_filtros["/api/planosmanutencao/query"]]
    assert "ativo" not in propriedades, "planos deve trazer ativos e inativos (tela cadastral)"


def test_planos_das_duas_unidades_usam_o_mesmo_criterio(app, criar_usuarios, capturar_filtros):
    resultados = {}
    for login, prefixo in (("ana.hmb", "/api/brasilandia"), ("beto.hetrin", "/api/grandmassif")):
        c = app.test_client()
        fazer_login(c, login, criar_usuarios)
        c.get(f"{prefixo}/planos")
        resultados[prefixo] = {f["property"] for f in capturar_filtros["/api/planosmanutencao/query"]}
    assert resultados["/api/brasilandia"] == resultados["/api/grandmassif"]


def test_exportacao_de_planos_valida_o_corpo(client, criar_usuarios):
    fazer_login(client, "ana.hmb", criar_usuarios)
    assert client.post("/api/brasilandia/exportar_planos", json={}).status_code == 400
    assert client.post("/api/brasilandia/exportar_planos", json={"itens": "x"}).status_code == 400
    ok = client.post("/api/brasilandia/exportar_planos", json={"itens": [{"id": 1, "descricao": "=1+1", "ativo": True}]})
    assert ok.status_code == 200
    from io import BytesIO
    from openpyxl import load_workbook
    celulas = [c.value for row in load_workbook(BytesIO(ok.data)).active.iter_rows(min_row=3) for c in row]
    assert "'=1+1" in celulas or "=1+1" not in celulas   # fórmula neutralizada ou ausente (nunca literal executável)
