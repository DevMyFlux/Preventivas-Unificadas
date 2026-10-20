"""Excel das preventivas: OS atrasadas saem com status "<status> em atraso" e a coluna "Período em atraso" diz há quanto tempo."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from core.exporters import descrever_atraso, gerar_excel_preventivas
from tests.test_api_unidade import *  # noqa: F401,F403  (fixtures do módulo)
from tests.test_api_unidade import P, _massa

HOJE = date(2026, 10, 14)


@pytest.mark.parametrize("prevista,esperado", [
    (date(2026, 10, 13), "1 dia"),
    (date(2026, 10, 10), "4 dias"),
    (date(2026, 9, 20), "24 dias"),               # menos de um mês de calendário → dias
    (date(2026, 9, 14), "1 mês"),
    (date(2026, 9, 10), "1 mês e 4 dias"),
    (date(2026, 8, 14), "2 meses"),
    (date(2026, 7, 1), "3 meses e 13 dias"),
    (date(2025, 10, 14), "12 meses"),
    (date(2026, 10, 14), ""),                     # hoje não é atraso
    (date(2026, 10, 15), ""),                     # futuro também não
])
def test_descrever_atraso_em_dias_ou_meses(prevista, esperado):
    assert descrever_atraso(prevista, HOJE) == esperado


def test_descrever_atraso_fim_de_mes_nao_estoura():
    assert descrever_atraso(date(2026, 1, 31), date(2026, 3, 1)) == "1 mês e 1 dia"     # 31/01 + 1 mês = 28/02
    assert descrever_atraso(date(2026, 1, 31), date(2026, 2, 28)) == "28 dias"


def _linhas(r):
    ws = load_workbook(io.BytesIO(r.data)).active
    cab = [c.value for c in ws[2]]
    return cab, [l for l in ws.iter_rows(min_row=3, values_only=True) if any(l)]


def test_excel_atrasada_vira_status_em_atraso_com_periodo(ana, neovero):
    _massa(neovero)
    r = ana.get(f"{P}/exportar_preventivas", query_string={"data_ini": "2026-10-14", "data_fim": "2026-10-14",
                                                           "status": "Pendente", "atrasada": "1"})
    cab, linhas = _linhas(r)
    assert cab[3] == "Status" and cab[4] == "Período em atraso" and "Atrasada" not in cab
    assert {(l[0], l[3], l[4]) for l in linhas} == {("10/10/2026", "Pendente em atraso", "4 dias"),
                                                    ("20/09/2026", "Pendente em atraso", "24 dias")}


def test_excel_so_a_linha_atrasada_muda_as_demais_ficam_como_estao(ana, neovero):
    _massa(neovero)
    r = ana.get(f"{P}/exportar_preventivas", query_string={"data_ini": "2026-10-14", "data_fim": "2026-10-14", "status": "Pendente"})
    _, linhas = _linhas(r)
    por_data = {l[0]: (l[3], l[4]) for l in linhas}
    assert por_data["14/10/2026"] == ("Pendente", None)                       # de hoje: não é atraso
    assert por_data["10/10/2026"] == ("Pendente em atraso", "4 dias")
    r = ana.get(f"{P}/exportar_preventivas", query_string={"data_ini": "2026-10-01", "data_fim": "2026-10-31", "status": "Concluída"})
    assert all(l[3] == "Concluída" and l[4] is None for l in _linhas(r)[1])   # concluída nunca atrasa


def test_excel_usa_a_data_de_hoje_informada():
    linha = {"data_iso": "2026-09-10", "data_prev": "10/09/2026", "atrasada": True, "status": "Em andamento", "plano": "X"}
    ws = load_workbook(gerar_excel_preventivas([linha], hoje=date(2026, 10, 14))).active
    assert [c.value for c in ws[3]][3:5] == ["Em andamento em atraso", "1 mês e 4 dias"]
