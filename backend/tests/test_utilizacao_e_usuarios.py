"""
Análise de utilização (somente Dev): tempo de sessão × ativo × ocioso, estado agora, usuários sem
utilização, critérios explícitos — e gestão de usuários: ficha, desbloqueio, perfil imutável,
proteção da conta administrativa e auditoria das operações.
"""
import json
from datetime import timedelta

import pytest
from sqlalchemy import select

from core import relatorios_uso
from core.auth import servico
from core.db import agora_utc, sessao
from core.models import AuditLog, AuthSession, User
from tests.conftest import fazer_login


@pytest.fixture()
def dev(app, criar_usuarios):
    c = app.test_client()
    fazer_login(c, "dev", criar_usuarios)
    return c


def _sessao_de(username: str) -> AuthSession:
    with sessao() as s:
        uid = s.scalar(select(User.id).where(User.username == username))
        return s.scalars(select(AuthSession).where(AuthSession.user_id == uid).order_by(AuthSession.id)).all()[-1]


def _ajustar(username: str, **campos):
    with sessao() as s:
        uid = s.scalar(select(User.id).where(User.username == username))
        sess = s.scalars(select(AuthSession).where(AuthSession.user_id == uid).order_by(AuthSession.id)).all()[-1]
        for k, v in campos.items():
            setattr(sess, k, v)


def _uso(dev, **q):
    return dev.get("/api/admin/usage", query_string=q).get_json()


# ── tempo de sessão × ativo × ocioso ─────────────────────────────────────────

def test_tempo_de_sessao_e_diferente_do_tempo_ativo(app, dev, criar_usuarios):
    ana = app.test_client()
    fazer_login(ana, "ana.hmb", criar_usuarios)
    agora = agora_utc()
    # sessão de 60 min no total, mas só 12 min de atividade efetiva
    _ajustar("ana.hmb", created_at=agora - timedelta(minutes=60), last_activity_at=agora - timedelta(minutes=1),
             active_seconds=12 * 60, activity_count=9)
    item = {i["username"]: i for i in _uso(dev)["itens"]}["ana.hmb"]
    assert item["tempo_sessao_min"] == pytest.approx(60, abs=1.5)      # sessão aberta, ainda dentro do limite
    assert item["tempo_ativo_min"] == pytest.approx(12, abs=0.2)
    assert item["tempo_ocioso_min"] == pytest.approx(item["tempo_sessao_min"] - 12, abs=0.2)
    assert item["acessos"] == 1


def test_estado_agora_ativo_ocioso_e_offline(app, dev, criar_usuarios):
    ana, beto = app.test_client(), app.test_client()
    fazer_login(ana, "ana.hmb", criar_usuarios)
    fazer_login(beto, "beto.hetrin", criar_usuarios)
    _ajustar("beto.hetrin", last_activity_at=agora_utc() - timedelta(minutes=12))       # aberta, sem atividade recente
    por = {i["username"]: i for i in _uso(dev)["itens"]}
    assert por["ana.hmb"]["estado_agora"] == "ativo"
    assert por["beto.hetrin"]["estado_agora"] == "ocioso"
    assert por["dev"]["estado_agora"] == "ativo"
    ind = _uso(dev)["indicadores"]
    assert ind["usuarios_ativos_agora"] == 2 and ind["usuarios_ociosos_agora"] == 1
    agora_lista = {s["username"]: s["estado"] for s in _uso(dev)["ativos_agora"]}
    assert agora_lista == {"ana.hmb": "ativo", "beto.hetrin": "ocioso", "dev": "ativo"}
    ana.post("/api/auth/logout")
    assert {i["username"]: i for i in _uso(dev)["itens"]}["ana.hmb"]["estado_agora"] == "offline"


def test_sessao_vencida_nao_aparece_como_ativa_agora_e_conta_como_encerrada_por_inatividade(app, dev, criar_usuarios):
    ana = app.test_client()
    fazer_login(ana, "ana.hmb", criar_usuarios)
    _ajustar("ana.hmb", last_activity_at=agora_utc() - timedelta(minutes=40))
    dados = _uso(dev)
    por = {i["username"]: i for i in dados["itens"]}
    assert por["ana.hmb"]["estado_agora"] == "offline"
    assert por["ana.hmb"]["sessoes_por_inatividade"] == 1
    assert "ana.hmb" not in {s["username"] for s in dados["ativos_agora"]}
    assert dados["indicadores"]["sessoes_por_inatividade"] == 1
    sessoes = dev.get("/api/admin/sessions").get_json()["itens"]
    da_ana = next(s for s in sessoes if s["username"] == "ana.hmb")
    assert da_ana["encerramento"] == "idle_timeout" and da_ana["por_inatividade"] is True
    assert da_ana["motivo_texto"] == "Expirou por inatividade" and da_ana["fim"]


def test_historico_de_sessoes_traz_inicio_fim_motivo_tempos_e_unidade(app, dev, criar_usuarios):
    ana = app.test_client()
    fazer_login(ana, "ana.hmb", criar_usuarios)
    ana.post("/api/auth/logout")
    s = next(x for x in dev.get("/api/admin/sessions").get_json()["itens"] if x["username"] == "ana.hmb")
    for campo in ("inicio", "fim", "encerramento", "motivo_texto", "duracao_min", "tempo_ativo_min", "tempo_ocioso_min",
                  "atividades", "unit", "estado", "ultima_atividade"):
        assert campo in s
    assert s["encerramento"] == "logout" and s["unit"] == "HMB" and s["motivo_texto"] == "Saiu (logout)"


# ── usuários sem utilização: critérios explícitos ────────────────────────────

def test_usuario_sem_registro_e_marcado_como_sem_utilizacao_registrada_com_parecer_e_nunca_como_nao_utiliza(dev):
    d = _uso(dev)
    beto = {i["username"]: i for i in d["itens"]}["beto.hetrin"]
    assert beto["utilizou"] is False and beto["situacao"] == "Sem utilização registrada"
    assert beto["acessos"] == 0 and beto["tempo_sessao_min"] == 0
    assert "não utiliza" not in json.dumps(d, ensure_ascii=False).lower().replace("nao utiliza", "")
    assert d["indicadores"]["usuarios_sem_utilizacao_registrada"] >= 2          # ana e beto
    assert any("não comprova" in c for c in d["criterios"])
    assert d["limite_inatividade_min"] == 25


def test_parecer_distingue_periodo_coberto_de_periodo_anterior_a_coleta(dev):
    hoje = agora_utc().date().isoformat()
    coberto = {i["username"]: i for i in _uso(dev, data_ini=hoje, data_fim=hoje)["itens"]}["beto.hetrin"]
    assert "período coberto" in coberto["parecer"]
    antigo = {i["username"]: i for i in _uso(dev, data_ini="2020-01-01", data_fim="2020-01-31")["itens"]}["beto.hetrin"]
    assert "a coleta de utilização começou em" in antigo["parecer"] and "não é coberto" in antigo["parecer"]


def test_filtros_por_unidade_usuario_e_periodo(app, dev, criar_usuarios):
    ana, beto = app.test_client(), app.test_client()
    fazer_login(ana, "ana.hmb", criar_usuarios)
    fazer_login(beto, "beto.hetrin", criar_usuarios)
    so_hmb = {i["username"] for i in _uso(dev, unit="HMB")["itens"]}
    assert so_hmb == {"ana.hmb"}
    uid = servico.obter_usuario_por_username("beto.hetrin").id
    assert {i["username"] for i in _uso(dev, user_id=uid)["itens"]} == {"beto.hetrin"}
    assert all(i["acessos"] == 0 for i in _uso(dev, data_ini="2020-01-01", data_fim="2020-01-02")["itens"])


def test_indicadores_batem_com_a_tabela_por_usuario(app, dev, criar_usuarios):
    for u in ("ana.hmb", "beto.hetrin"):
        fazer_login(app.test_client(), u, criar_usuarios)
    d = _uso(dev)
    itens = d["itens"]
    ind = d["indicadores"]
    assert ind["acessos"] == sum(i["acessos"] for i in itens)
    assert ind["tempo_ativo_min"] == pytest.approx(sum(i["tempo_ativo_min"] for i in itens), abs=0.2)
    assert ind["usuarios_que_utilizaram"] == sum(1 for i in itens if i["utilizou"])
    assert ind["usuarios_cadastrados"] == len(itens) == 3


def test_usuario_desativado_continua_no_relatorio_com_todo_o_historico(app, dev, criar_usuarios):
    ana = app.test_client()
    fazer_login(ana, "ana.hmb", criar_usuarios)
    ana.get("/api/brasilandia/habilidades")
    uid = servico.obter_usuario_por_username("ana.hmb").id
    dev.post(f"/api/admin/users/{uid}/active", json={"active": False})
    item = {i["username"]: i for i in _uso(dev)["itens"]}["ana.hmb"]
    assert item["conta_ativa"] is False and item["utilizou"] is True and item["acessos"] == 1
    assert item["estado_agora"] == "offline"
    ficha = dev.get(f"/api/admin/users/{uid}").get_json()
    assert ficha["usuario"]["is_active"] is False and ficha["sessoes"] and ficha["eventos"]


def test_exportacao_de_uso_traz_as_colunas_de_tempo_e_os_criterios(dev):
    from io import BytesIO
    from openpyxl import load_workbook
    r = dev.get("/api/admin/usage/export")
    ws = load_workbook(BytesIO(r.data)).active
    cab = [c.value for c in ws[2]]
    for coluna in ("Tempo de sessão (min)", "Tempo ativo (min)", "Tempo ocioso (min)", "Sessões encerradas por inatividade", "Parecer"):
        assert coluna in cab
    textos = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "não comprova" in textos


# ── somente o Dev consulta (no backend) ──────────────────────────────────────

@pytest.mark.parametrize("rota", ["/api/admin/usage", "/api/admin/usage/export", "/api/admin/sessions", "/api/admin/audit",
                                  "/api/admin/audit/summary", "/api/admin/audit/export", "/api/admin/users/1"])
def test_dados_globais_de_utilizacao_so_para_o_dev(app, criar_usuarios, rota):
    ana = app.test_client()
    fazer_login(ana, "ana.hmb", criar_usuarios)
    assert ana.get(rota).status_code == 403
    assert app.test_client().get(rota).status_code == 401


# ── ficha da conta, desbloqueio, perfis ──────────────────────────────────────

def test_ficha_do_usuario_traz_cadastro_utilizacao_sessoes_e_eventos(app, dev, criar_usuarios):
    ana = app.test_client()
    fazer_login(ana, "ana.hmb", criar_usuarios)
    ana.post("/api/auth/activity", json={"module": "planos", "unit": "HMB"})
    uid = servico.obter_usuario_por_username("ana.hmb").id
    f = dev.get(f"/api/admin/users/{uid}").get_json()
    assert f["usuario"]["username"] == "ana.hmb" and f["usuario"]["unit"] == "HMB" and f["usuario"]["sessoes_abertas"] == 1
    assert f["utilizacao"]["acessos"] == 1 and f["utilizacao"]["estado_agora"] == "ativo"
    assert f["sessoes"][0]["estado"] == "ativo"
    assert any(e["event"] == "module_view" for e in f["eventos"]) and f["total_eventos"] >= 2
    assert "password" not in json.dumps(f).lower().replace("must_change_password", "")
    assert dev.get("/api/admin/users/9999").status_code == 404


def test_desbloquear_usuario_limpa_o_bloqueio_temporario_e_audita(app, dev, criar_usuarios, monkeypatch):
    monkeypatch.setenv("LOGIN_MAX_FAILURES", "2")
    c = app.test_client()
    for _ in range(2):
        c.post("/api/auth/login", json={"username": "ana.hmb", "password": "errada-errada1"})
    assert c.post("/api/auth/login", json={"username": "ana.hmb", "password": criar_usuarios}).status_code == 423
    uid = servico.obter_usuario_por_username("ana.hmb").id
    assert dev.get("/api/admin/users").get_json()["itens"][-1] is not None
    assert dev.post(f"/api/admin/users/{uid}/unlock").status_code == 200
    assert c.post("/api/auth/login", json={"username": "ana.hmb", "password": criar_usuarios}).status_code == 200
    with sessao() as s:
        ev = s.scalars(select(AuditLog).where(AuditLog.path == f"/api/admin/users/{uid}/unlock")).all()
        assert ev and ev[-1].username == "dev" and ev[-1].result == "success"


def test_perfil_nao_pode_ser_alterado_nem_promovido_por_ninguem(app, dev, criar_usuarios):
    uid = servico.obter_usuario_por_username("ana.hmb").id
    for corpo in ({"role": "dev"}, {"is_dev": True}, {"role": "user"}):
        r = dev.patch(f"/api/admin/users/{uid}", json=corpo)
        assert r.status_code == 400 and r.get_json()["codigo"] == "perfil_imutavel"
    assert servico.obter_usuario_por_username("ana.hmb").role == "user"
    ana = app.test_client()
    fazer_login(ana, "ana.hmb", criar_usuarios)
    assert ana.patch(f"/api/admin/users/{uid}", json={"role": "dev"}).status_code == 403     # usuário comum nem chega lá
    assert ana.post("/api/admin/users", json={"username": "x@y.com", "password": "Senha12345", "role": "dev"}).status_code == 403
    assert servico.obter_usuario_por_username("ana.hmb").role == "user"


def test_usuario_comum_nao_altera_a_propria_unidade(app, criar_usuarios):
    ana = app.test_client()
    fazer_login(ana, "ana.hmb", criar_usuarios)
    uid = servico.obter_usuario_por_username("ana.hmb").id
    assert ana.patch(f"/api/admin/users/{uid}", json={"unit": "HETRIN"}).status_code == 403
    assert servico.obter_usuario_por_username("ana.hmb").unit == "HMB"
    assert ana.patch("/api/auth/me", json={"unit": "HETRIN"}).status_code in (404, 405)


def test_a_conta_administrativa_nao_pode_ser_desativada_nem_ficar_sem_administrador(app, dev):
    admin_id = servico.obter_usuario_por_username("dev").id
    r = dev.post(f"/api/admin/users/{admin_id}/active", json={"active": False})
    assert r.status_code == 400 and r.get_json()["codigo"] in ("admin_proprio", "dev_protegido")
    assert dev.get("/api/auth/me").status_code == 200          # continua logado
    assert servico.existe_dev()
    # mesmo mexendo no vínculo ou nome, o Dev segue sem unidade (acessa todas)
    assert dev.patch(f"/api/admin/users/{admin_id}", json={"unit": "HMB"}).status_code == 400
    assert servico.obter_usuario_por_username("dev").unit is None


def test_operacoes_administrativas_ficam_na_auditoria_com_detalhes_sem_segredos(app, dev, criar_usuarios):
    dev.post("/api/admin/users", json={"username": "novo.teste", "password": "SenhaSecreta123", "display_name": "Novo", "unit": "HMB"})
    uid = servico.obter_usuario_por_username("novo.teste").id
    dev.patch(f"/api/admin/users/{uid}", json={"unit": "HETRIN"})
    dev.post(f"/api/admin/users/{uid}/active", json={"active": False})
    dev.post(f"/api/admin/users/{uid}/active", json={"active": True})
    dev.post(f"/api/admin/users/{uid}/password", json={"password": "OutraSenha456"})
    with sessao() as s:
        ev = s.scalars(select(AuditLog).where(AuditLog.module == "admin").order_by(AuditLog.id)).all()
    caminhos = [e.path for e in ev if e.method in ("POST", "PATCH")]
    for esperado in ("/api/admin/users", f"/api/admin/users/{uid}", f"/api/admin/users/{uid}/active", f"/api/admin/users/{uid}/password"):
        assert esperado in caminhos
    tudo = " ".join(str(e.details) for e in ev)
    assert "SenhaSecreta123" not in tudo and "OutraSenha456" not in tudo
    assert all(e.username == "dev" for e in ev)


def test_criar_usuario_exige_unidade_valida_e_impede_duplicidade_de_login_com_email(dev):
    ok = dev.post("/api/admin/users", json={"username": "Fulano@Empresa.com", "password": "Senha12345", "unit": "HETRIN"})
    assert ok.status_code == 201 and ok.get_json()["username"] == "fulano@empresa.com"
    dup = dev.post("/api/admin/users", json={"username": "fulano@empresa.com", "password": "Senha12345", "unit": "HMB"})
    assert dup.status_code == 400 and dup.get_json()["campo"] == "username"


def test_nota_metodologica_e_criterios_nao_afirmam_inatividade_por_falta_de_registro():
    texto = " ".join(relatorios_uso.criterios())
    assert "não comprova" in texto and "NÃO significa que a pessoa estava usando" in texto
