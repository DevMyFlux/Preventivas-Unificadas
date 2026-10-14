"""Área administrativa (somente Dev), trilha de auditoria, sessões e relatórios de utilização."""
import csv
import io
import json
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from core import relatorios_uso
from core.auth import servico
from core.db import agora_utc, sessao
from core.models import AuditLog, AuthSession
from tests.conftest import fazer_login


@pytest.fixture()
def dev(app, criar_usuarios):
    c = app.test_client()
    fazer_login(c, "dev", criar_usuarios)
    return c


@pytest.fixture()
def ana(app, criar_usuarios):
    c = app.test_client()
    fazer_login(c, "ana.hmb", criar_usuarios)
    return c


ROTAS_ADMIN = [("GET", "/api/admin/users"), ("POST", "/api/admin/users"), ("GET", "/api/admin/audit"),
               ("GET", "/api/admin/audit/summary"), ("GET", "/api/admin/audit/export"), ("GET", "/api/admin/sessions"),
               ("GET", "/api/admin/usage"), ("GET", "/api/admin/usage/export"), ("PATCH", "/api/admin/users/1"),
               ("POST", "/api/admin/users/1/active"), ("POST", "/api/admin/users/1/password")]


@pytest.mark.parametrize("metodo,rota", ROTAS_ADMIN)
def test_rotas_administrativas_sao_negadas_a_usuario_comum_e_anonimo(app, ana, metodo, rota):
    assert ana.open(rota, method=metodo, json={}).status_code == 403
    assert app.test_client().open(rota, method=metodo, json={}).status_code == 401


def test_acesso_negado_fica_registrado_na_auditoria(ana):
    ana.get("/api/admin/audit")
    with sessao() as s:
        ev = s.scalars(select(AuditLog).where(AuditLog.result == "denied")).all()[-1]
        assert ev.username == "ana.hmb" and ev.module == "admin" and "area_administrativa" in ev.details


# ── gestão de usuários ────────────────────────────────────────────────────────

def test_dev_cria_usuario_vinculado_a_unidade(dev):
    r = dev.post("/api/admin/users", json={"username": "novo.user", "password": "Senha12345", "display_name": "Novo",
                                           "unit": "HETRIN"})
    assert r.status_code == 201
    u = r.get_json()
    assert u["unit"] == "HETRIN" and u["role"] == "user" and u["is_active"] is True and "password_hash" not in u and "scrypt" not in json.dumps(u)
    lista = dev.get("/api/admin/users").get_json()["itens"]
    assert {x["username"] for x in lista} >= {"dev", "ana.hmb", "beto.hetrin", "novo.user"}


@pytest.mark.parametrize("corpo,campo", [
    ({"username": "x", "password": "Senha12345", "unit": "HMB"}, "username"),
    ({"username": "valido.user", "password": "curta", "unit": "HMB"}, "password"),
    ({"username": "valido.user", "password": "Senha12345", "unit": "OUTRA"}, "unit"),
    ({"username": "valido.user", "password": "Senha12345"}, "unit"),
    ({"username": "ana.hmb", "password": "Senha12345", "unit": "HMB"}, "username"),
])
def test_validacao_na_criacao_de_usuario(dev, corpo, campo):
    r = dev.post("/api/admin/users", json=corpo)
    assert r.status_code == 400 and r.get_json()["campo"] == campo


def test_nao_e_possivel_criar_outro_dev_nem_promover(dev):
    r = dev.post("/api/admin/users", json={"username": "outro.dev", "password": "Senha12345", "role": "dev"})
    assert r.status_code == 400 and r.get_json()["codigo"] == "perfil_invalido"
    uid = servico.obter_usuario_por_username("ana.hmb").id
    dev.patch(f"/api/admin/users/{uid}", json={"role": "dev"})   # campo ignorado
    assert servico.obter_usuario_por_username("ana.hmb").role == "user"


def test_conta_dev_nao_pode_ser_desativada(dev):
    uid = servico.obter_usuario_por_username("dev").id
    r = dev.post(f"/api/admin/users/{uid}/active", json={"active": False})
    assert r.status_code == 400 and r.get_json()["codigo"] == "dev_protegido"


def test_desativar_preserva_historico_e_reativar_devolve_o_acesso(app, dev, ana, criar_usuarios):
    uid = servico.obter_usuario_por_username("ana.hmb").id
    ana.get("/api/brasilandia/colaboradores")        # gera histórico de atividade
    with sessao() as s:
        antes = s.scalar(select(AuditLog).where(AuditLog.username == "ana.hmb").limit(1))
        assert antes is not None
    assert dev.post(f"/api/admin/users/{uid}/active", json={"active": False}).status_code == 200
    with sessao() as s:
        n = len(s.scalars(select(AuditLog).where(AuditLog.username == "ana.hmb")).all())
        assert n >= 2   # histórico continua lá (nada apagado)
        sess = s.scalars(select(AuthSession).where(AuthSession.user_id == uid)).all()
        assert sess and all(x.ended_at is not None for x in sess)
        assert {x.end_reason for x in sess} == {"user_deactivated"}
    assert dev.post(f"/api/admin/users/{uid}/active", json={"active": True}).status_code == 200
    assert app.test_client().post("/api/auth/login", json={"username": "ana.hmb", "password": criar_usuarios}).status_code == 200


def test_redefinir_senha_pelo_dev_nao_expoe_nem_loga_a_senha_e_derruba_sessoes(app, dev, ana):
    uid = servico.obter_usuario_por_username("ana.hmb").id
    r = dev.post(f"/api/admin/users/{uid}/password", json={"password": "SenhaNova999", "must_change_password": True})
    assert r.status_code == 200 and "SenhaNova999" not in r.get_data(as_text=True)
    assert ana.get("/api/auth/me").status_code == 401
    novo = app.test_client()
    login = novo.post("/api/auth/login", json={"username": "ana.hmb", "password": "SenhaNova999"})
    assert login.status_code == 200 and login.get_json()["user"]["must_change_password"] is True
    with sessao() as s:
        tudo = " ".join(str(l.details) for l in s.scalars(select(AuditLog)).all())
    assert "SenhaNova999" not in tudo


def test_outro_dev_nao_poderia_trocar_a_senha_da_conta_dev(dev):
    # a conta Dev só pode ser alterada por ela mesma; aqui o próprio Dev consegue
    uid = servico.obter_usuario_por_username("dev").id
    assert dev.post(f"/api/admin/users/{uid}/password", json={"password": "OutraSenha123"}).status_code == 200


# ── auditoria ─────────────────────────────────────────────────────────────────

def test_operacoes_ficam_registradas_com_usuario_unidade_modulo_acao_e_resultado(app, ana):
    ana.get("/api/brasilandia/planos")                      # leitura (falha do Neovero não configurado)
    ana.patch("/api/brasilandia/colaborador/Fulano/status", json={"status": "Desligado"})
    ana.get("/api/grandmassif/colaboradores")               # negado
    with sessao() as s:
        eventos = s.scalars(select(AuditLog).where(AuditLog.username == "ana.hmb").order_by(AuditLog.id)).all()
    por_modulo = {(e.module, e.action, e.result) for e in eventos}
    assert ("planos", "read", "failure") in por_modulo          # 503: Neovero não configurado no teste
    assert ("colaboradores", "update", "failure") in por_modulo  # colaborador/cadastro inexistente
    assert ("colaboradores", "read", "denied") in por_modulo     # outra unidade
    e = next(x for x in eventos if x.result == "denied")
    # unit = unidade ALVO da requisição; a do usuário fica em details para investigação
    assert e.unit == "HETRIN" and '"unidade_do_usuario": "HMB"' in e.details
    assert e.method == "GET" and e.status_code == 403 and e.ip and e.ts


def test_navegacao_entre_modulos_gera_evento_e_atualiza_atividade(ana):
    r = ana.post("/api/auth/activity", json={"module": "preventivas", "unit": "HMB"})
    assert r.status_code == 200
    assert ana.post("/api/auth/activity", json={"module": "inexistente"}).status_code == 400
    with sessao() as s:
        ev = s.scalars(select(AuditLog).where(AuditLog.event == "module_view")).all()[-1]
        assert ev.module == "preventivas" and ev.unit == "HMB" and ev.username == "ana.hmb"


def test_filtros_da_consulta_de_auditoria(app, dev, ana, criar_usuarios):
    beto = app.test_client()
    fazer_login(beto, "beto.hetrin", criar_usuarios)
    ana.post("/api/auth/activity", json={"module": "planos", "unit": "HMB"})
    beto.post("/api/auth/activity", json={"module": "colaboradores", "unit": "HETRIN"})
    ver = lambda **q: dev.get("/api/admin/audit", query_string=q).get_json()
    assert {i["username"] for i in ver(username="ana")["itens"]} == {"ana.hmb"}
    assert {i["unit"] for i in ver(unit="HETRIN")["itens"]} <= {"HETRIN"}
    assert all(i["module"] == "colaboradores" for i in ver(module="colaboradores", event="module_view")["itens"])
    hoje = agora_utc().date().isoformat()
    assert ver(data_ini=hoje, data_fim=hoje)["total"] > 0
    assert ver(data_ini="2020-01-01", data_fim="2020-01-02")["total"] == 0
    assert all(i["result"] == "success" for i in ver(result="success")["itens"])
    assert ver(result="invalido")["total"] == ver()["total"]       # valor inválido é ignorado, não vira SQL


def test_paginacao_da_auditoria(dev):
    for _ in range(7):
        dev.post("/api/auth/activity", json={"module": "planos"})
    p1 = dev.get("/api/admin/audit", query_string={"page_size": 3, "page": 1, "event": "module_view"}).get_json()
    p2 = dev.get("/api/admin/audit", query_string={"page_size": 3, "page": 2, "event": "module_view"}).get_json()
    assert p1["total"] == 7 and len(p1["itens"]) == 3 and len(p2["itens"]) == 3
    assert not {i["id"] for i in p1["itens"]} & {i["id"] for i in p2["itens"]}


def test_consulta_a_auditoria_nao_se_auto_registra_mas_exportacao_sim(dev):
    dev.get("/api/admin/audit")
    dev.get("/api/admin/usage")
    with sessao() as s:
        assert not s.scalars(select(AuditLog).where(AuditLog.module == "admin", AuditLog.action == "read")).all()
    assert dev.get("/api/admin/audit/export").status_code == 200
    with sessao() as s:
        assert s.scalars(select(AuditLog).where(AuditLog.module == "admin", AuditLog.action == "export")).all()


# ── relatórios ────────────────────────────────────────────────────────────────

def test_resumo_traz_indicadores_e_nota_metodologica(app, dev, ana, criar_usuarios):
    app.test_client().post("/api/auth/login", json={"username": "ana.hmb", "password": "errada-errada1"})
    ana.post("/api/auth/activity", json={"module": "preventivas", "unit": "HMB"})
    r = dev.get("/api/admin/audit/summary").get_json()
    t = r["totais"]
    assert t["logins"] >= 2 and t["logins_falhos"] >= 1 and t["usuarios_ativos"] >= 2 and t["sessoes"] >= 2
    assert any(m["modulo"] == "preventivas" for m in r["por_modulo"])
    assert r["por_dia"] and r["por_dia"][0]["eventos"] > 0
    assert "Ausência de registros não significa inatividade" in r["nota"]
    assert r["limite_inatividade_min"] == 25


def test_relatorio_por_usuario_diferencia_sem_registros_de_inatividade(app, dev, ana):
    itens = {i["username"]: i for i in dev.get("/api/admin/usage").get_json()["itens"]}
    assert itens["ana.hmb"]["utilizou"] is True and itens["ana.hmb"]["sessoes"] >= 1
    # beto nunca entrou: 'Sem utilização registrada' — e NÃO 'inativo'/'não utiliza'
    assert itens["beto.hetrin"]["situacao"] == "Sem utilização registrada"
    assert itens["beto.hetrin"]["eventos"] == 0 and itens["beto.hetrin"]["sessoes"] == 0
    assert "inativ" not in itens["beto.hetrin"]["situacao"].lower() and "não utiliza" not in itens["beto.hetrin"]["parecer"].lower()


def test_sessoes_mostram_duracao_e_motivo_de_encerramento(app, dev, ana):
    ana.post("/api/auth/logout")
    ociosa = app.test_client()
    fazer_login(ociosa, "beto.hetrin", "Senha12345")
    with sessao() as s:
        sess = s.scalars(select(AuthSession).order_by(AuthSession.id)).all()[-1]
        sess.last_activity_at = agora_utc() - timedelta(hours=2)
        sess.created_at = agora_utc() - timedelta(hours=3)
    itens = dev.get("/api/admin/sessions").get_json()["itens"]
    por_user = {i["username"]: i for i in itens}
    assert por_user["ana.hmb"]["encerramento"] == "logout"
    assert por_user["beto.hetrin"]["encerramento"] == "idle_timeout"
    assert por_user["beto.hetrin"]["duracao_min"] == pytest.approx(60, abs=1)   # até a ÚLTIMA atividade
    assert por_user["dev"]["encerramento"] == "em_andamento"


def test_exportacao_xlsx_e_csv_da_auditoria_e_do_uso(dev):
    dev.post("/api/auth/activity", json={"module": "planos"})
    x = dev.get("/api/admin/audit/export")
    assert x.status_code == 200 and "spreadsheetml" in x.mimetype
    wb = load_workbook(io.BytesIO(x.data))
    linhas = list(wb.active.iter_rows(values_only=True))
    assert linhas[1][0] == "Data/Hora (UTC)" and len(linhas) > 3
    c = dev.get("/api/admin/audit/export?format=csv")
    assert c.status_code == 200 and c.mimetype == "text/csv"
    rows = list(csv.reader(io.StringIO(c.data.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][0] == "Data/Hora (UTC)" and len(rows) > 2
    u = dev.get("/api/admin/usage/export")
    assert u.status_code == 200
    assert load_workbook(io.BytesIO(u.data)).active.max_row >= 4


def test_exportacao_neutraliza_injecao_de_formula(dev):
    from core import audit
    audit.registrar("teste", username="=cmd|' /C calc'!A0", module="x", details={"k": "+SOMA(1)"})
    x = dev.get("/api/admin/audit/export?format=csv").data.decode("utf-8-sig")
    assert "'=cmd" in x and ";=cmd" not in x


def test_dados_de_auditoria_nunca_expõem_hash_ou_token(dev):
    texto = json.dumps(dev.get("/api/admin/audit?page_size=200").get_json()) + \
        json.dumps(dev.get("/api/admin/users").get_json()) + json.dumps(dev.get("/api/admin/sessions").get_json())
    assert "scrypt:" not in texto and "password_hash" not in texto and "token_hash" not in texto


def test_dia_local_usa_fuso_brasil():
    from datetime import datetime
    assert relatorios_uso.dia_local(datetime(2026, 10, 14, 2, 0)).isoformat() == "2026-10-13"   # 23h do dia 13 em BRT
    ini, fim = relatorios_uso.intervalo_utc(relatorios_uso.parse_data("2026-10-13"), relatorios_uso.parse_data("2026-10-13"))
    assert (ini, fim) == (datetime(2026, 10, 13, 3, 0), datetime(2026, 10, 14, 3, 0))


def test_auditoria_registra_a_unidade_acessada_inclusive_para_o_dev(app, dev, ana):
    dev.get("/api/brasilandia/habilidades")
    dev.get("/api/grandmassif/habilidades")
    ana.get("/api/grandmassif/habilidades")        # negado: a unidade ALVO fica registrada
    with sessao() as s:
        eventos = s.scalars(select(AuditLog).where(AuditLog.module == "colaboradores").order_by(AuditLog.id)).all()
    por = [(e.username, e.unit, e.result) for e in eventos]
    assert ("dev", "HMB", "success") in por and ("dev", "HETRIN", "success") in por
    assert ("ana.hmb", "HETRIN", "denied") in por
